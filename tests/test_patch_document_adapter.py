from __future__ import annotations

import importlib.util
import json
import sys
import types
import unittest
import zipfile
from pathlib import Path
from unittest import mock
from uuid import uuid4


MODULE_PATH = Path(__file__).resolve().parents[1] / "scripts" / "patch_document_adapter.py"
SPEC = importlib.util.spec_from_file_location("patch_document_adapter", MODULE_PATH)
assert SPEC and SPEC.loader
adapter = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = adapter
SPEC.loader.exec_module(adapter)


class FakePage:
    def __init__(self, text: str) -> None:
        self.text = text
        self.calls = 0

    def extract_text(self, **options: int) -> str:
        self.calls += 1
        return self.text


class FakePdf:
    def __init__(self, pages: list[FakePage]) -> None:
        self.pages = pages

    def __enter__(self) -> "FakePdf":
        return self

    def __exit__(self, *arguments: object) -> None:
        return None


class PatchDocumentAdapterTest(unittest.TestCase):
    def temporary_file(self, suffix: str) -> Path:
        path = MODULE_PATH.parents[1] / f".patch-adapter-{uuid4().hex}{suffix}"
        self.addCleanup(lambda: path.unlink(missing_ok=True))
        return path

    def test_normalization_and_generic_html_preserve_parseable_headings(self) -> None:
        self.assertEqual(adapter.normalize_text("\ufeffAlpha\u200b Beta"), "Alpha Beta")
        rendered = adapter.strip_html(
            "<h1>Patch Notes</h1><p>Body</p><h2>Items</h2><p>Changed.</p>"
        )
        self.assertIn("# Patch Notes", rendered)
        self.assertIn("## Items", rendered)

    def test_html_extraction_enforces_the_text_ceiling(self) -> None:
        source = self.temporary_file(".html")
        source.write_text("<h1>August 22, 2026</h1><p>Long historical body.</p>", encoding="utf-8")

        result = adapter.extract_html(source, max_text_characters=8)

        self.assertEqual(result["status"], "limit_exceeded")
        self.assertEqual(result["limit"], "max_text_characters")

    def test_feed_reports_raw_accepted_rejected_and_truncated_items(self) -> None:
        source = self.temporary_file(".rss")
        source.write_text(
            """<?xml version="1.0" encoding="UTF-8"?>
<rss><channel>
  <item><title>Accepted</title><description>Body</description></item>
  <item><description>No title</description></item>
  <item><title>No body</title></item>
  <item><title>Truncated</title><description>Later</description></item>
</channel></rss>""",
            encoding="utf-8",
        )
        result = adapter.extract_feed(source, max_items=3)

        self.assertEqual(result["status"], "feed")
        self.assertEqual(result["raw_item_count"], 4)
        self.assertEqual(result["processed_item_count"], 3)
        self.assertEqual(result["accepted_item_count"], 1)
        self.assertEqual(result["rejected_item_count"], 2)
        self.assertEqual(result["truncated_item_count"], 1)
        self.assertEqual(result["item_count"], 1)
        self.assertEqual(result["records"][0]["item_index"], 1)
        self.assertEqual(result["rejections"], [
            {"item_index": 2, "reasons": ["missing_title"]},
            {"item_index": 3, "reasons": ["missing_content_element"]},
        ])

    def test_json_structured_records_map_aliases_and_preserve_item_indices(self) -> None:
        source = self.temporary_file(".json")
        source.write_text(json.dumps([
            {
                "name": "Game Update Notes: August 22, 2026",
                "body": "A sufficiently long structured patch body.",
                "date": "2026-08-22",
                "link": "https://forums.everquest.com/index.php?threads/example.1/",
                "excerpt": True,
            },
            "not an object",
            {"title": "August 23, 2026"},
            {
                "title": "August 24, 2026",
                "name": "A conflicting title",
                "content": "Another sufficiently long patch body.",
            },
        ]), encoding="utf-8")

        result = adapter.extract_json_records(source, max_items=20)

        self.assertEqual(result["status"], "structured_records")
        self.assertEqual(
            (result["raw_item_count"], result["processed_item_count"],
             result["accepted_item_count"], result["rejected_item_count"],
             result["truncated_item_count"]),
            (4, 4, 1, 3, 0),
        )
        self.assertEqual(result["records"][0]["item_index"], 1)
        self.assertEqual(result["records"][0]["title"], "Game Update Notes: August 22, 2026")
        self.assertEqual(result["records"][0]["content"], "A sufficiently long structured patch body.")
        self.assertTrue(result["records"][0]["excerpt"])
        self.assertEqual([item["item_index"] for item in result["rejections"]], [2, 3, 4])
        self.assertEqual(result["rejections"][0]["reasons"], ["record_not_object"])
        self.assertEqual(result["rejections"][1]["reasons"], ["missing_content"])
        self.assertEqual(result["rejections"][2]["reasons"], ["conflicting_title_fields"])

    def test_json_object_collections_and_malformed_shapes_are_bounded(self) -> None:
        source = self.temporary_file(".json")
        source.write_text(json.dumps({"patches": [
            {"display_date": "August 22, 2026", "notes": "A sufficiently long patch body."},
            {"display_date": "August 23, 2026", "notes": "A second sufficiently long patch body."},
        ]}), encoding="utf-8")
        truncated = adapter.extract_json_records(source, max_items=1)
        self.assertEqual(
            (truncated["raw_item_count"], truncated["accepted_item_count"], truncated["truncated_item_count"]),
            (2, 1, 1),
        )

        limited = adapter.extract_json_records(source, max_items=20, max_text_characters=10)
        self.assertEqual((limited["status"], limited["limit"]), ("limit_exceeded", "max_text_characters"))
        self.assertEqual(limited["records"], [])

        source.write_text('{"patches": [], "records": []}', encoding="utf-8")
        ambiguous = adapter.extract_json_records(source, max_items=20)
        self.assertEqual(ambiguous["reason_code"], "ambiguous_or_missing_record_array")

        source.write_text('{"patches": [], "patches": []}', encoding="utf-8")
        duplicate = adapter.extract_json_records(source, max_items=20)
        self.assertEqual(duplicate["reason_code"], "malformed_json")

        source.write_text('[{"title":"August 22, 2026","content":NaN}]', encoding="utf-8")
        nonfinite = adapter.extract_json_records(source, max_items=20)
        self.assertEqual(nonfinite["reason_code"], "malformed_json")

    def test_json_unsafe_url_is_omitted_with_indexed_warning(self) -> None:
        source = self.temporary_file(".json")
        source.write_text(json.dumps({"records": [{
            "title": "August 22, 2026",
            "content": "A sufficiently long patch body.",
            "source_url": "https://archive-user:do-not-leak@example.com/private",
        }]}), encoding="utf-8")
        result = adapter.extract_json_records(source, max_items=20)
        self.assertEqual(result["accepted_item_count"], 1)
        self.assertIsNone(result["records"][0]["source_url"])
        self.assertEqual(result["warnings"][0]["item_index"], 1)
        self.assertEqual(result["warnings"][0]["code"], "unsafe_source_url")
        self.assertNotIn("do-not-leak", json.dumps(result))

    def test_archive_json_and_csv_prefer_display_date_but_generic_aliases_remain_strict(self) -> None:
        patch = {
            "id": "2026-08-22-1",
            "slug": "2026-08-22-1",
            "patch_date": "2026-08-22",
            "display_date": "Game Update Notes: August 22, 2026",
            "title": "August 22, 2026",
            "content": "A sufficiently long exported patch body.",
            "content_hash": "a" * 64,
            "source_occurrences": [],
        }
        json_source = self.temporary_file(".json")
        json_source.write_text(json.dumps({
            "schema_version": 1,
            "record_count": 1,
            "patches": [patch],
        }), encoding="utf-8")
        json_result = adapter.extract_json_records(json_source, max_items=20)
        self.assertEqual(json_result["accepted_item_count"], 1)
        self.assertEqual(json_result["records"][0]["title"], patch["display_date"])

        csv_source = self.temporary_file(".csv")
        csv_source.write_text(
            "id,slug,patch_date,effective_date,display_date,title,sequence,content,content_hash,source_occurrences\n"
            "2026-08-23-1,2026-08-23-1,2026-08-23,2026-08-23,"
            '"Game Update Notes: August 23, 2026","August 23, 2026",1,'
            '"Another sufficiently long exported patch body., with punctuation",'
            f'{"b" * 64},[]\n',
            encoding="utf-8",
        )
        csv_result = adapter.extract_delimited_records(csv_source, "csv", ",", 20)
        self.assertEqual(csv_result["accepted_item_count"], 1)
        self.assertEqual(csv_result["records"][0]["title"], "Game Update Notes: August 23, 2026")

        generic_source = self.temporary_file(".json")
        generic_source.write_text(json.dumps({"records": [{
            "display_date": "Game Update Notes: August 24, 2026",
            "title": "August 24, 2026",
            "content": "A sufficiently long generic patch body.",
        }]}), encoding="utf-8")
        generic_result = adapter.extract_json_records(generic_source, max_items=20)
        self.assertEqual(generic_result["accepted_item_count"], 0)
        self.assertEqual(generic_result["rejections"][0]["reasons"], ["conflicting_title_fields"])

    def test_csv_and_tsv_headers_rows_and_reason_codes(self) -> None:
        csv_source = self.temporary_file(".csv")
        csv_source.write_text(
            "display_date,notes,patch_date,url,excerpt\n"
            '"August 22, 2026","A sufficiently long CSV patch body.",2026-08-22,https://example.com/patch,true\n'
            '"August 23, 2026","Another sufficiently long patch body.",2026-08-23,https://example.com/patch,false,extra\n',
            encoding="utf-8",
        )
        csv_result = adapter.extract_delimited_records(csv_source, "csv", ",", 20)
        self.assertEqual(
            (csv_result["raw_item_count"], csv_result["accepted_item_count"], csv_result["rejected_item_count"]),
            (2, 1, 1),
        )
        self.assertEqual((csv_result["records"][0]["item_index"], csv_result["records"][0]["row_index"]), (2, 2))
        self.assertEqual(csv_result["rejections"][0]["item_index"], 3)
        self.assertEqual(csv_result["rejections"][0]["reasons"], ["column_count_mismatch"])

        tsv_source = self.temporary_file(".tsv")
        tsv_source.write_text(
            "name\tbody\tdate\n"
            "August 24, 2026\tA sufficiently long TSV patch body.\t2026-08-24\n",
            encoding="utf-8",
        )
        tsv_result = adapter.extract_delimited_records(tsv_source, "tsv", "\t", 20)
        self.assertEqual(tsv_result["accepted_item_count"], 1)
        self.assertEqual(tsv_result["records"][0]["source_kind"], "tsv_record")

        tsv_source.write_text("title\tTITLE\tcontent\nA\tA\tLong enough content\n", encoding="utf-8")
        duplicate_header = adapter.extract_delimited_records(tsv_source, "tsv", "\t", 20)
        self.assertEqual(duplicate_header["reason_code"], "duplicate_header")

    def test_csv_field_limit_tracks_config_and_accepts_fields_larger_than_128_kib(self) -> None:
        source = self.temporary_file(".csv")
        body = "x" * (140 * 1024)
        source.write_text(f'title,content\n"August 25, 2026","{body}"\n', encoding="utf-8")
        configured_limit = len(body) + 1_024

        with mock.patch.object(adapter.csv, "field_size_limit", wraps=adapter.csv.field_size_limit) as field_limit:
            result = adapter.extract_delimited_records(
                source,
                "csv",
                ",",
                max_items=20,
                max_text_characters=configured_limit,
            )

        field_limit.assert_called_once_with(configured_limit)
        self.assertEqual(result["accepted_item_count"], 1)
        self.assertEqual(len(result["records"][0]["content"]), len(body))

    def test_utf16_dtd_is_rejected_before_feed_parsing(self) -> None:
        source = self.temporary_file(".rss")
        source.write_bytes((
            '<?xml version="1.0" encoding="utf-16"?>'
            '<!DOCTYPE rss [<!ENTITY x "unsafe">]>'
            '<rss><channel><item><title>&x;</title><description>Body</description></item>'
            '</channel></rss>'
        ).encode("utf-16"))

        result = adapter.extract_feed(source, max_items=20)
        self.assertEqual(result["status"], "unsafe_xml_rejected")
        self.assertEqual(result["raw_item_count"], None)
        self.assertEqual(result["records"], [])

        declaration = '<!DOCTYPE rss><rss><channel /></rss>'
        for encoding in ("utf-16-le", "utf-16-be"):
            with self.subTest(encoding=encoding), self.assertRaises(ValueError):
                adapter.reject_unsafe_xml_declarations(declaration.encode(encoding))

    def test_feed_byte_limit_prevents_xml_parsing(self) -> None:
        source = self.temporary_file(".rss")
        source.write_bytes(b"<rss><channel /></rss>")
        result = adapter.extract_feed(source, max_items=20, max_xml_bytes=8)
        self.assertEqual(result["status"], "limit_exceeded")
        self.assertEqual(result["limit"], "max_xml_bytes")

    def test_pdf_page_limit_is_checked_before_page_extraction(self) -> None:
        pages = [FakePage("one"), FakePage("two"), FakePage("three")]
        fake_pdfplumber = types.SimpleNamespace(open=lambda source: FakePdf(pages))
        with mock.patch.dict(sys.modules, {"pdfplumber": fake_pdfplumber}):
            result = adapter.extract_pdf(Path("unused.pdf"), 1, max_pages=2)
        self.assertEqual(result["status"], "limit_exceeded")
        self.assertEqual(result["limit"], "max_pdf_pages")
        self.assertEqual([page.calls for page in pages], [0, 0, 0])

    def test_pdf_text_limit_stops_bounded_extraction(self) -> None:
        pages = [FakePage("abcdefghij"), FakePage("klmnopqrst")]
        fake_pdfplumber = types.SimpleNamespace(open=lambda source: FakePdf(pages))
        with mock.patch.dict(sys.modules, {"pdfplumber": fake_pdfplumber}):
            result = adapter.extract_pdf(
                Path("unused.pdf"),
                1,
                max_pages=10,
                max_text_characters=15,
            )
        self.assertEqual(result["status"], "limit_exceeded")
        self.assertEqual(result["limit"], "max_text_characters")
        self.assertEqual([page.calls for page in pages], [1, 1])
        self.assertEqual(result["text"], "")

    def write_zip(
        self,
        entries: list[tuple[str, bytes]],
        *,
        compression: int = zipfile.ZIP_DEFLATED,
    ) -> Path:
        source = self.temporary_file(".docx")
        with zipfile.ZipFile(source, "w", compression=compression) as archive:
            for name, value in entries:
                archive.writestr(name, value)
        return source

    def test_docx_zip_entry_and_size_limits_are_structured(self) -> None:
        entry_source = self.write_zip([
            ("one.bin", b"1"),
            ("two.bin", b"2"),
            ("three.bin", b"3"),
        ])
        result = adapter.extract_docx(entry_source, 1, max_entries=2)
        self.assertEqual((result["status"], result["limit"]), ("limit_exceeded", "max_docx_entries"))

        uncompressed_source = self.write_zip([("word/document.xml", b"A" * 100)])
        result = adapter.extract_docx(uncompressed_source, 1, max_uncompressed_bytes=50)
        self.assertEqual(result["limit"], "max_docx_uncompressed_bytes")

        entry_size_source = self.write_zip([("word/document.xml", b"A" * 100)])
        result = adapter.extract_docx(entry_size_source, 1, max_entry_uncompressed_bytes=50)
        self.assertEqual(result["limit"], "max_docx_entry_uncompressed_bytes")

        compressed_source = self.write_zip([("stored.bin", b"A" * 100)], compression=zipfile.ZIP_STORED)
        result = adapter.extract_docx(compressed_source, 1, max_compressed_bytes=20)
        self.assertEqual(result["limit"], "max_docx_compressed_bytes")

    def test_docx_compression_ratio_and_utf16_dtd_are_rejected(self) -> None:
        ratio_source = self.write_zip([("word/document.xml", b"A" * 10_000)])
        result = adapter.extract_docx(ratio_source, 1, max_compression_ratio=2)
        self.assertEqual(result["limit"], "max_compression_ratio")

        unsafe_xml = (
            '<?xml version="1.0" encoding="utf-16"?>'
            '<!DOCTYPE document [<!ENTITY x "unsafe">]><document>&x;</document>'
        ).encode("utf-16")
        unsafe_source = self.write_zip([("word/document.xml", unsafe_xml)])
        result = adapter.extract_docx(unsafe_source, 1)
        self.assertEqual(result["status"], "unsafe_xml_rejected")
        self.assertEqual(result["entry"], "word/document.xml")

    def test_normal_docx_preserves_heading_and_enforces_paragraph_and_text_limits(self) -> None:
        from docx import Document

        source = self.temporary_file(".docx")
        document = Document()
        document.add_heading("Patch Notes", level=1)
        document.add_paragraph("Body\u200b text")
        document.save(source)

        result = adapter.extract_docx(source, 1)
        self.assertEqual(result["status"], "text")
        self.assertIn("# Patch Notes", result["text"])
        self.assertIn("Body text", result["text"])
        self.assertNotIn("\u200b", result["text"])

        paragraph_limited = adapter.extract_docx(source, 1, max_paragraphs=1)
        self.assertEqual(paragraph_limited["limit"], "max_docx_paragraphs")

        text_limited = adapter.extract_docx(source, 1, max_text_characters=5)
        self.assertEqual(text_limited["limit"], "max_text_characters")

    def test_stdout_is_forced_to_utf8(self) -> None:
        class Console:
            def __init__(self) -> None:
                self.configuration = None

            def reconfigure(self, **options: str) -> None:
                self.configuration = options

        console = Console()
        adapter.configure_utf8_stdout(console)
        self.assertEqual(console.configuration, {"encoding": "utf-8", "errors": "strict"})


if __name__ == "__main__":
    unittest.main()
